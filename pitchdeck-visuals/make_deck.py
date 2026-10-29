"""PowerPoint pitch deck (16:9) following the Intel AI Global Impact Festival judging flow:
problem & impact -> solution -> AI & Intel tech -> data & model -> results -> demo -> responsible AI -> scale.

  inference/.venv/Scripts/python pitchdeck-visuals/make_deck.py
Logo boxes are placeholders: paste the official Intel / festival logos there (not redrawn here).
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
S = HERE / "slides"
V = HERE / "videos"
BG, PANEL, TXT, MUTED, YEL, BLUE, RED = (11, 13, 16), (21, 24, 29), (244, 244, 245), (161, 161, 170), (255, 214, 10), (76, 141, 255), (226, 30, 40)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(c):
    return RGBColor(*c)


def box(sl, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    return s


def text(sl, x, y, w, h, t, size=18, color=TXT, bold=False):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = rgb(color); p.font.name = "Segoe UI"
    return tb


def header(sl, section):
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = rgb(BG)
    for i, lab in enumerate(["LOGO INTEL", "LOGO INTEL AI GLOBAL\nIMPACT FESTIVAL 2026"]):
        b = box(sl, 0.4 + i * 1.9, 0.25, 1.75 if i == 0 else 2.4, 0.55, None, (90, 90, 100))
        b.text_frame.text = lab
        for p in b.text_frame.paragraphs:
            p.font.size = Pt(8); p.font.color.rgb = rgb(MUTED)
    box(sl, 12.25, 0.3, 0.7, 0.22, (227, 38, 54), shape=MSO_SHAPE.RECTANGLE)  # Indonesia flag
    box(sl, 12.25, 0.52, 0.7, 0.22, (255, 255, 255), shape=MSO_SHAPE.RECTANGLE)
    text(sl, 7.2, 0.3, 4.9, 0.5, section, 12, MUTED)


def content(section, title, bullets, icon="●", image=None):
    sl = prs.slides.add_slide(BLANK)
    header(sl, section)
    text(sl, 0.6, 1.1, 12, 0.9, title, 34, TXT, True)
    y = 2.2
    w = 6.0 if image else 12
    for b in bullets:
        box(sl, 0.6, y + 0.08, 0.42, 0.42, YEL, shape=MSO_SHAPE.OVAL).text_frame.text = icon
        text(sl, 1.2, y, w - 0.6, 0.9, b, 17)
        y += 0.95
    if image and Path(image).exists():
        sl.shapes.add_picture(str(image), Inches(6.9), Inches(2.0), width=Inches(6.0))
    return sl


def full(image, video=None):
    sl = prs.slides.add_slide(BLANK)
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = rgb(BG)
    sl.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
    if video and Path(video).exists():
        sl.shapes.add_movie(str(video), Inches(9.6), Inches(4.6), Inches(3.4), Inches(1.9), mime_type="video/mp4")
    return sl


def video_slide(title, vid):
    sl = prs.slides.add_slide(BLANK)
    header(sl, "Demo")
    text(sl, 0.6, 1.1, 12, 0.8, title, 30, TXT, True)
    if vid.exists():
        sl.shapes.add_movie(str(vid), Inches(1.9), Inches(2.0), Inches(9.5), Inches(5.3), mime_type="video/mp4")


# 1. Title
sl = prs.slides.add_slide(BLANK); header(sl, "Intel AI Global Impact Festival 2026 · Indonesia")
text(sl, 0.6, 2.2, 12, 1.2, "PAVEMENT-PINN", 60, TXT, True)
text(sl, 0.6, 3.5, 12, 1, "Potret jalan rusak. AI dan petugas menanganinya sampai tuntas.", 26, YEL)
text(sl, 0.6, 4.4, 12, 1, "Computer Vision + Physics-Informed Neural Network di atas Intel® OpenVINO™", 20, MUTED)
# 2-3. Problem & impact
content("1 · Latar belakang masalah", "Jalan rusak: bahaya nyata, penanganan lambat", [
    "Lubang & genangan membahayakan pengendara motor, terutama saat hujan dan malam hari.",
    "Laporan warga tersebar (media sosial, telepon) tanpa bukti terstruktur dan sulit dilacak.",
    "Petugas harus turun ke lokasi hanya untuk memastikan jenis dan ukuran kerusakan.",
    "Perbaikan sering tambal-sulam: penyebab struktural (tanah dasar, truk ODOL, musim hujan) tidak dihitung.",
    "Warga tidak tahu laporannya ditangani atau tidak, sehingga kepercayaan turun."], "!")
content("2 · Dampak & SDG", "Siapa yang terbantu", [
    "Warga: lapor dalam 30 detik, pantau sampai selesai, Rute Aman & Radar Bahaya.",
    "Petugas: inspeksi terpandu 3 sisi, ukur di foto, rekomendasi berbasis fisika.",
    "Koordinator & pimpinan: antrian prioritas, SLA, verifikasi AI, rapor instansi.",
    "SDG 9 (infrastruktur tangguh), SDG 11 (kota aman), SDG 3.6 (keselamatan jalan), SDG 16 (institusi akuntabel)."], "✓")
# 4. Solution
content("3 · Solusi", "Satu alur, empat peran, semua tercatat", [
    "Warga memotret 3 sisi → AI mengenali 11 jenis kerusakan + genangan.",
    "Koordinator menerima & menugaskan → petugas inspeksi terpandu + ukur di foto.",
    "PINN menghitung sisa umur dengan lapisan jalan, beban truk, dan curah hujan Open-Meteo.",
    "SPK → perbaikan → AI memeriksa foto sesudah → warga mengonfirmasi."], "→", S / "1_model_dan_data" / "15-15_alur_laporan.png")
# 5. Intel tech
content("4 · Teknologi AI & Intel", "Semua model berjalan lokal di Intel® OpenVINO™", [
    "YOLOv8n deteksi + YOLOv8s-seg segmentasi retak + CLIP ViT-B/32 scene (INT8, NNCF).",
    "PINN surrogate (FP32/FP16/INT8) + solver MLET PyMastic sebagai rujukan fisika.",
    "SLM Qwen2.5-1.5B INT4 via OpenVINO GenAI untuk draf usulan (angka dari mesin).",
    "Siap CPU, iGPU, dan NPU Intel Core Ultra (bench/run_all.py)."], "AI", S / "1_model_dan_data" / "14-14_arsitektur_sistem.png")
# 6+. Data & model, results (chart slides)
for f in sorted((S / "1_model_dan_data").glob("[0-9][0-9]-*.png")):
    if f.name.startswith(("14-", "15-", "16-")):
        continue
    full(f)
video_slide("Training PINN (rekaman run asli)", V / "training_pinn.mp4")
video_slide("Proses deteksi AI pada foto jalan nyata", V / "deteksi_jalan-berlubang.mp4")
# Demo: both simulations
for sim, vids in (("2_simulasi_A", ["simulasi_A_warga-melapor.mp4", "simulasi_A_petugas-inspeksi.mp4"]), ("3_simulasi_B", ["simulasi_B_warga-melapor.mp4", "simulasi_B_petugas-inspeksi.mp4"])):
    files = sorted((S / sim).glob("*.png"))
    for f in files:
        full(f)
    for v in vids:
        video_slide(f"Rekaman layar: {v.replace('.mp4', '').replace('_', ' ')}", V / v)
# Responsible AI, feasibility, roadmap, closing
content("6 · AI yang bertanggung jawab", "Manusia tetap memutuskan", [
    "Setiap keputusan staf wajib beralasan dan tercatat (audit log).",
    "Hasil AI dari klien dibuang; server hanya memakai analisis dari penyimpanannya sendiri.",
    "EXIF foto dihapus, lokasi dari GPS saat itu, bahasa netral tanpa tuduhan.",
    "Pengusul penanganan tidak bisa menyetujui sendiri; angka diberi label terukur/asumsi."], "⚖")
content("7 · Kelayakan & rencana", "Dari pilot ke skala kota", [
    "Pilot: DKI Jakarta + Karawang dengan 5 instansi (nasional, provinsi, kota/kabupaten).",
    "Berjalan di laptop biasa; siap VPS/Docker (deploy_vm.sh, docker-compose).",
    "Berikutnya: uji NPU Intel Core Ultra, akurasi lapangan ≥ 20 foto/kelas, uji SUS warga & petugas.",
    "Skala: basis data, integrasi SIM jalan instansi, data WIM untuk beban truk."], "↗")
sl = prs.slides.add_slide(BLANK); header(sl, "Terima kasih")
text(sl, 0.6, 2.5, 12, 1.2, "Terima kasih", 54, TXT, True)
text(sl, 0.6, 3.8, 12, 1, "Coba langsung: https://avlfelo.tail8b73d2.ts.net", 24, YEL)
text(sl, 0.6, 4.6, 12, 1, "Kode & dokumentasi: github.com/muhammadkiro516-netizen/Pavement-PINN-", 18, MUTED)
out = HERE / "PAVEMENT-PINN_pitchdeck_final.pptx"
prs.save(out)
print("wrote", out, len(prs.slides._sldIdLst), "slides")
